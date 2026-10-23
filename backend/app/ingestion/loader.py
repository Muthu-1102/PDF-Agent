"""
Document loaders. Each function returns page-aware text where possible,
so upstream chunking can retain accurate page_number metadata.
"""

from typing import List, Tuple

import fitz  # PyMuPDF
import docx


def load_pdf(file_path: str) -> List[Tuple[int, str]]:
    """Returns list of (page_number, page_text)."""
    pages = []
    with fitz.open(file_path) as doc:
        for i, page in enumerate(doc, start=1):
            pages.append((i, page.get_text()))
    return pages


def load_docx(file_path: str) -> List[Tuple[int, str]]:
    """DOCX has no native page concept -> treat whole doc as page 1.
    H10 Fix: Extract both paragraphs and table rows so structured data is not lost."""
    document = docx.Document(file_path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_cells:
                lines.append(" | ".join(row_cells))
    return [(1, "\n".join(lines))]


def load_txt(file_path: str) -> List[Tuple[int, str]]:
    """H10 Fix: Handle non-UTF8 encodings cleanly without raising 500."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return [(1, f.read())]
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            return [(1, f.read())]


def load_document(file_path: str, file_type: str) -> List[Tuple[int, str]]:
    file_type = file_type.lower()
    if file_type == "pdf":
        return load_pdf(file_path)
    if file_type == "docx":
        return load_docx(file_path)
    if file_type == "txt":
        return load_txt(file_path)
    raise ValueError(f"Unsupported file type: {file_type}")
